"""Word (.docx) versions of the fictional sample statements.

Uses the same figures as generate_sample_statements.py, so the PDF and Word
versions always agree. Each page carries a diagonal watermark reading
"SAMPLE - NOT A REAL STATEMENT".
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor

from generate_sample_statements import (
    ACCOUNTS, BANK, HOLDER, MONTHS, OPENING, OUT_DIR, YEAR, compute_month, money,
)

WATERMARK_XML = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
    '<w:pict>'
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e">'
    '<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/>'
    '<v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/>'
    '<v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/>'
    '<v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/>'
    '<v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom"/>'
    '<v:textpath on="t" fitshape="t"/><o:lock v:ext="edit" text="t" shapetype="t"/>'
    '</v:shapetype>'
    '<v:shape id="SampleWatermark" o:spid="_x0000_s1025" type="#_x0000_t136" '
    'style="position:absolute;margin-left:0;margin-top:0;width:520pt;height:70pt;rotation:315;'
    'z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
    'mso-position-vertical:center;mso-position-vertical-relative:margin" '
    'fillcolor="#E06666" stroked="f">'
    '<v:fill opacity=".35"/>'
    '<v:textpath on="t" fitshape="t" style="font-family:&quot;Arial&quot;;font-size:1pt" '
    'string="SAMPLE - NOT A REAL STATEMENT"/>'
    '</v:shape></w:pict></w:r>'
)


def para(doc_or_cell, text="", bold=False, size=10, align=None, color=None):
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def table(doc, header, rows, right_cols=(), widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = ""
        para(t.rows[0].cells[i], h, bold=True, size=9)
        t.rows[0].cells[i].paragraphs[0]._p.getparent().remove(t.rows[0].cells[i].paragraphs[0]._p)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].text = ""
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(9)
            if isinstance(val, str) and val.startswith("Total"):
                r.bold = True
            if i in right_cols:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i in right_cols:
        t.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        t.autofit = False
        for col, w in zip(t.columns, widths):
            col.width = Inches(w)
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)
    return t


def build_docx(m):
    doc = Document()
    sec = doc.sections[0]
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(m["stmt_date"])
    hr.font.size = Pt(9)
    hp._p.append(parse_xml(WATERMARK_XML))
    para(sec.header, "SAMPLE - NOT A REAL STATEMENT", bold=True, size=11,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0xC0, 0, 0))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Fictional test data for control-testing purposes only. "
                    "Not issued by any financial institution.")
    fr.italic, fr.font.size, fr.font.color.rgb = True, Pt(7), RGBColor(0xC0, 0, 0)

    # Page 1: summary
    para(doc, BANK, bold=True, size=16)
    para(doc, ACCOUNTS["checking"][0], bold=True, size=12)
    for line in HOLDER:
        para(doc, line)
    para(doc, "Questions? 1-000-000-0000 (fictional)  |  Online: example.invalid", size=9)
    para(doc)
    para(doc, "Accounts linked to your primary account", bold=True, size=11)
    rows = []
    for key in ("checking", "everyday", "savings"):
        label, num = ACCOUNTS[key]
        suffix = " - primary account" if key == "checking" else ""
        rows.append((f"{label} ({num}){suffix}", f"{m['balances'][key]:,.2f}"))
    rows.append(("Total qualification balance this month", money(m["total"])))
    table(doc, ("Account (Account Number)", "$ Balance"), rows, right_cols=(1,), widths=(4.8, 1.7))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Page 2: activity
    month, last = m["month"], m["last"]
    para(doc, f"{ACCOUNTS['checking'][0]}  -  {ACCOUNTS['checking'][1]}", bold=True, size=12)
    para(doc, "Statement period activity", bold=True, size=11)
    table(doc, ("Item", "Amount"), [
        (f"Balance on {month}/1", money(m["opening"])),
        ("Deposits/Additions", money(m["deposits"])),
        ("Withdrawals/Subtractions", f"- {money(m['withdrawals'])}"),
        (f"Balance on {month}/{last}", money(m["closing"])),
        ("Interest paid this statement", money(m["interest"])),
    ], right_cols=(1,), widths=(4.8, 1.7))
    para(doc)
    para(doc, "Transaction history", bold=True, size=11)
    tx = [(d, desc, f"{dep:,.2f}" if dep else "", f"{wd:,.2f}" if wd else "", f"{run:,.2f}")
          for d, desc, dep, wd, run in m["rows"]]
    tx.append(("Totals", "", money(m["deposits"]), money(m["withdrawals"]), ""))
    table(doc, ("Date", "Description", "Deposits/Additions", "Withdrawals", "Ending Balance"),
          tx, right_cols=(2, 3, 4), widths=(0.6, 3.0, 1.0, 0.95, 1.0))

    doc.core_properties.title = f"SAMPLE statement {m['name']} {YEAR} - fictional test data"
    path = OUT_DIR / f"SAMPLE_Statement_{m['name']}{YEAR}.docx"
    doc.save(path)
    print(f"{path.name}: close {money(m['closing'])}")


if __name__ == "__main__":
    bal = dict(OPENING)
    for month in MONTHS:
        m = compute_month(month, bal)
        build_docx(m)
        bal = m["balances"]
